"""DOCX 解析器 —— 基于 python-docx 提取文本、表格，保留标题层级。"""

from app.parser.document_parser import DocumentParser


class DocxParser(DocumentParser):
    """Word 文档解析器：提取段落文本和表格内容，保留 Heading 样式用于结构标记。"""

    def extract_text(self, file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            parts = []

            # 按文档元素顺序遍历（段落 + 表格交替出现）
            for element in doc.element.body:
                tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag

                if tag == 'p':
                    # 段落元素
                    para = self._find_paragraph(doc, element)
                    if para:
                        text = para.text.strip()
                        if text and len(text) > 3:
                            if para.style.name.startswith('Heading'):
                                parts.append(f"\n{'=' * 40}\n{text}\n{'=' * 40}\n")
                            else:
                                parts.append(text)

                elif tag == 'tbl':
                    # 表格元素
                    table = self._find_table(doc, element)
                    if table:
                        table_text = self._extract_table(table)
                        if table_text:
                            parts.append(table_text)

            return "\n\n".join(parts)

        except ImportError:
            return f"[无法解析的 DOCX 内容: {file_path}]"

    def _find_paragraph(self, doc, element):
        """根据 XML 元素查找对应的 python-docx Paragraph 对象。"""
        for para in doc.paragraphs:
            if para._element is element:
                return para
        return None

    def _find_table(self, doc, element):
        """根据 XML 元素查找对应的 python-docx Table 对象。"""
        for table in doc.tables:
            if table._element is element:
                return table
        return None

    def _extract_table(self, table) -> str:
        """将 Word 表格转为带结构的文本，保留行列关系。"""
        rows = table.rows
        if not rows:
            return ""

        lines = ["【表格】"]
        max_cols = max(len(row.cells) for row in rows)

        for row_idx, row in enumerate(rows):
            cells = []
            for cell in row.cells:
                cell_text = cell.text.strip().replace('\n', ' ')
                cells.append(cell_text)

            # 补齐列数不足的行
            while len(cells) < max_cols:
                cells.append("")

            line = " | ".join(cells)
            if row_idx == 0:
                # 表头行加分隔线
                lines.append(line)
                lines.append("-" * min(len(line), 120))
            else:
                lines.append(line)

        return "\n".join(lines)
